"""Build an editable PowerPoint deck: title, summary, QC, one slide per figure."""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt

WIDE = (Inches(13.333), Inches(7.5))


def _title(slide, text, size=28):
    box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12.3), Inches(0.9))
    p = box.text_frame.paragraphs[0]
    p.text, p.font.size, p.font.bold = text, Pt(size), True


def _bullets(slide, lines, top=1.3, size=16):
    box = slide.shapes.add_textbox(Inches(0.6), Inches(top), Inches(12), Inches(5.8))
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text, p.font.size = f"• {line}", Pt(size)


def build_deck(path: Path | str, title: str, subtitle: str, summary: list[str],
               qc_lines: list[str], figures: list[tuple[str, Path, str]]) -> Path:
    """figures = [(slide title, image path, caption)]"""
    prs = Presentation()
    prs.slide_width, prs.slide_height = WIDE
    blank = prs.slide_layouts[6]

    s = prs.slides.add_slide(blank)
    _title(s, title, 36)
    _bullets(s, [subtitle], top=1.5, size=18)

    s = prs.slides.add_slide(blank)
    _title(s, "Summary")
    _bullets(s, summary)

    s = prs.slides.add_slide(blank)
    _title(s, "Plate quality control")
    _bullets(s, qc_lines)

    for slide_title, img, caption in figures:
        s = prs.slides.add_slide(blank)
        _title(s, slide_title, 24)
        s.shapes.add_picture(str(img), Inches(0.6), Inches(1.2), height=Inches(5.3))
        cap = s.shapes.add_textbox(Inches(0.6), Inches(6.6), Inches(12), Inches(0.7)).text_frame
        cap.word_wrap = True
        cap.paragraphs[0].text, cap.paragraphs[0].font.size = caption, Pt(12)

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(path)
    return path
